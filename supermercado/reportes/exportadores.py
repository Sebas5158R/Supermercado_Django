import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from io import BytesIO
from datetime import date
from django.http import HttpResponse
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage, KeepTogether
from reportlab.lib.units import inch
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

AZUL_OSCURO = colors.HexColor('#1e293b')
AZUL_ACENTO = colors.HexColor('#38bdf8')
GRIS_CLARO = colors.HexColor('#f1f5f9')
BLANCO = colors.white

COLOR_FONDO = '#0f172a'
COLOR_BARRA = '#38bdf8'
COLORES_DONA = ['#38bdf8', '#34d399', '#f59e0b', '#f87171', '#a78bfa', '#fb923c']

plt.rcParams.update({
    'figure.facecolor': '#1e293b',
    'axes.facecolor': '#1e293b',
    'axes.edgecolor': '#334155',
    'axes.labelcolor': '#94a3b8',
    'xtick.color': '#64748b',
    'ytick.color': '#64748b',
    'text.color': '#D3DCED',
    'grid.color': '#334155',
    'grid.linestyle': '--',
    'grid.alpha': 0.5,
})


def _grafica_barras(etiquetas, valores, titulo, ancho=6.5, alto=3.0):
    fig, ax = plt.subplots(figsize=(ancho, alto))
    x = range(len(etiquetas))
    barras = ax.bar(x, valores, color=COLOR_BARRA, edgecolor='none', width=0.55)
    ax.set_xticks(x)
    etiqs_cortas = [str(e)[:18] + '…' if len(str(e)) > 18 else str(e) for e in etiquetas]
    ax.set_xticklabels(etiqs_cortas, rotation=30, ha='right', fontsize=8)
    ax.set_title(titulo, fontsize=10, fontweight='bold', color='#D3DCED', pad=10)
    ax.yaxis.grid(True)
    ax.set_axisbelow(True)
    for barra in barras:
        h = barra.get_height()
        if h > 0:
            ax.text(barra.get_x() + barra.get_width() / 2, h * 1.02,
                    f'{h:,.0f}', ha='center', va='bottom', fontsize=7, color='#38bdf8')
    fig.tight_layout()
    buf = BytesIO()
    fig.savefig(buf, format='png', dpi=130, bbox_inches='tight', facecolor='#1e293b')
    buf.seek(0)
    plt.close(fig)
    return buf


def _grafica_linea(etiquetas, valores, titulo, ancho=6.5, alto=2.8):
    fig, ax = plt.subplots(figsize=(ancho, alto))
    ax.plot(range(len(etiquetas)), valores, color=COLOR_BARRA, linewidth=2, marker='o', markersize=4)
    ax.fill_between(range(len(etiquetas)), valores, alpha=0.15, color=COLOR_BARRA)
    ax.set_xticks(range(len(etiquetas)))
    etiqs_cortas = [str(e)[5:] if len(str(e)) >= 10 else str(e) for e in etiquetas]
    ax.set_xticklabels(etiqs_cortas, rotation=30, ha='right', fontsize=7)
    ax.set_title(titulo, fontsize=10, fontweight='bold', color='#D3DCED', pad=10)
    ax.yaxis.grid(True)
    ax.set_axisbelow(True)
    fig.tight_layout()
    buf = BytesIO()
    fig.savefig(buf, format='png', dpi=130, bbox_inches='tight', facecolor='#1e293b')
    buf.seek(0)
    plt.close(fig)
    return buf


def _grafica_dona(etiquetas, valores, titulo, ancho=4.0, alto=3.2):
    valores_limpios = [max(v, 0) for v in valores]
    if sum(valores_limpios) == 0:
        valores_limpios = [1] * len(etiquetas)
    fig, ax = plt.subplots(figsize=(ancho, alto))
    wedges, texts, autotexts = ax.pie(
        valores_limpios,
        labels=None,
        colors=COLORES_DONA[:len(etiquetas)],
        autopct='%1.0f%%',
        startangle=90,
        wedgeprops={'edgecolor': '#1e293b', 'linewidth': 2},
        pctdistance=0.75,
    )
    for at in autotexts:
        at.set_fontsize(8)
        at.set_color('#0f172a')
        at.set_fontweight('bold')
    leyenda = [mpatches.Patch(color=COLORES_DONA[i % len(COLORES_DONA)], label=str(e)[:22])
               for i, e in enumerate(etiquetas)]
    ax.legend(handles=leyenda, loc='lower center', bbox_to_anchor=(0.5, -0.22),
              ncol=2, fontsize=7, framealpha=0, labelcolor='#94a3b8')
    centre = plt.Circle((0, 0), 0.5, fc='#1e293b')
    ax.add_patch(centre)
    ax.set_title(titulo, fontsize=10, fontweight='bold', color='#D3DCED', pad=10)
    fig.tight_layout()
    buf = BytesIO()
    fig.savefig(buf, format='png', dpi=130, bbox_inches='tight', facecolor='#1e293b')
    buf.seek(0)
    plt.close(fig)
    return buf


def _buf_a_rl_image(buf, ancho_inch, alto_inch):
    return RLImage(buf, width=ancho_inch * inch, height=alto_inch * inch)


def exportar_pdf(reporte, resultado):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=0.45 * inch,
        leftMargin=0.45 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch,
    )

    estilos = getSampleStyleSheet()
    est_titulo = ParagraphStyle('titulo', parent=estilos['Title'],
                                fontSize=17, textColor=AZUL_OSCURO, spaceAfter=2)
    est_sub = ParagraphStyle('sub', parent=estilos['Normal'],
                              fontSize=9, textColor=colors.HexColor('#64748b'), spaceAfter=4)
    est_seccion = ParagraphStyle('seccion', parent=estilos['Normal'],
                                  fontSize=11, textColor=AZUL_ACENTO, spaceAfter=8,
                                  fontName='Helvetica-Bold')
    est_celda = ParagraphStyle('celda', parent=estilos['Normal'], fontSize=8, wordWrap='CJK')
    est_kpi_et = ParagraphStyle('kpiet', parent=estilos['Normal'],
                                 fontSize=8, textColor=colors.HexColor('#94a3b8'))
    est_kpi_val = ParagraphStyle('kpival', parent=estilos['Normal'],
                                  fontSize=13, textColor=AZUL_ACENTO, fontName='Helvetica-Bold')

    contenido = []
    ancho_pagina = landscape(letter)[0] - 0.9 * inch

    contenido.append(Paragraph(reporte.nombre, est_titulo))
    contenido.append(Paragraph(
        f"Tipo: {reporte.get_tipo_display()}  ·  Generado: {date.today().strftime('%d/%m/%Y')}  ·  Registros: {resultado['cantidad']}",
        est_sub,
    ))
    if reporte.descripcion:
        contenido.append(Paragraph(reporte.descripcion, est_sub))
    contenido.append(Spacer(1, 0.12 * inch))

    totales = resultado.get('totales', {})
    if totales:
        contenido.append(Paragraph('Indicadores clave', est_seccion))
        items_kpi = list(totales.items())
        cols_kpi = min(len(items_kpi), 4)
        ancho_kpi = ancho_pagina / cols_kpi

        filas_kpi = []
        for i in range(0, len(items_kpi), cols_kpi):
            grupo = items_kpi[i:i + cols_kpi]
            while len(grupo) < cols_kpi:
                grupo.append(('', ''))
            fila_etq = [Paragraph(str(k), est_kpi_et) for k, v in grupo]
            fila_val = [Paragraph(str(v), est_kpi_val) for k, v in grupo]
            filas_kpi.extend([fila_etq, fila_val])

        tabla_kpi = Table(filas_kpi, colWidths=[ancho_kpi] * cols_kpi)
        tabla_kpi.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), AZUL_OSCURO),
            ('ROWBACKGROUNDS', (0, 0), (-1, -1), [colors.HexColor('#1e293b'), colors.HexColor('#1e293b')]),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('LEFTPADDING', (0, 0), (-1, -1), 14),
            ('LINEAFTER', (0, 0), (-2, -1), 0.5, colors.HexColor('#334155')),
            ('LINEBEFORE', (0, 0), (0, -1), 3, AZUL_ACENTO),
        ]))
        contenido.append(tabla_kpi)
        contenido.append(Spacer(1, 0.18 * inch))

    graficas = resultado.get('graficas', {})
    if graficas:
        contenido.append(Paragraph('Análisis gráfico', est_seccion))
        imgs_row = []
        for clave, datos in graficas.items():
            etqs = datos.get('etiquetas', [])
            vals = datos.get('valores', [])
            tipo = datos.get('tipo', 'bar')
            titulo_g = datos.get('titulo', clave)
            if not etqs or not vals:
                continue
            if tipo == 'line':
                buf = _grafica_linea(etqs, vals, titulo_g, ancho=6.0, alto=2.6)
                imgs_row.append(_buf_a_rl_image(buf, 4.8, 2.1))
            elif tipo == 'doughnut':
                buf = _grafica_dona(etqs, vals, titulo_g, ancho=3.8, alto=3.0)
                imgs_row.append(_buf_a_rl_image(buf, 3.2, 2.5))
            else:
                buf = _grafica_barras(etqs, vals, titulo_g, ancho=5.8, alto=2.6)
                imgs_row.append(_buf_a_rl_image(buf, 4.6, 2.1))

            if len(imgs_row) == 2:
                tabla_imgs = Table([imgs_row], colWidths=[ancho_pagina / 2] * 2)
                tabla_imgs.setStyle(TableStyle([
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 4),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ]))
                contenido.append(tabla_imgs)
                contenido.append(Spacer(1, 0.1 * inch))
                imgs_row = []

        if imgs_row:
            relleno = [Spacer(1, 1)] * (2 - len(imgs_row))
            tabla_imgs = Table([imgs_row + relleno], colWidths=[ancho_pagina / 2] * 2)
            tabla_imgs.setStyle(TableStyle([
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ]))
            contenido.append(tabla_imgs)
            contenido.append(Spacer(1, 0.1 * inch))

    filas = resultado.get('filas', [])
    etiquetas = resultado.get('etiquetas', {})
    if filas and etiquetas:
        contenido.append(Spacer(1, 0.1 * inch))
        contenido.append(Paragraph('Detalle de datos', est_seccion))
        claves = list(etiquetas.keys())
        encabezados = list(etiquetas.values())
        ancho_col = ancho_pagina / len(encabezados)
        datos_tabla = [encabezados]
        for fila in filas:
            fila_pdf = [Paragraph(str(fila.get(k, '')), est_celda) for k in claves]
            datos_tabla.append(fila_pdf)
        tabla = Table(datos_tabla, colWidths=[ancho_col] * len(encabezados), repeatRows=1)
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), AZUL_OSCURO),
            ('TEXTCOLOR', (0, 0), (-1, 0), AZUL_ACENTO),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [BLANCO, GRIS_CLARO]),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#cbd5e1')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        contenido.append(tabla)

    doc.build(contenido)
    buffer.seek(0)
    respuesta = HttpResponse(buffer, content_type='application/pdf')
    nombre_archivo = reporte.nombre.replace(' ', '_').lower()
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre_archivo}_{date.today()}.pdf"'
    return respuesta


def exportar_excel(reporte, resultado):
    libro = openpyxl.Workbook()

    hoja_dash = libro.active
    hoja_dash.title = 'Dashboard'

    hoja_datos = libro.create_sheet(title='Datos')

    relleno_oscuro = PatternFill(start_color='0f172a', end_color='0f172a', fill_type='solid')
    relleno_medio = PatternFill(start_color='1e293b', end_color='1e293b', fill_type='solid')
    relleno_acento = PatternFill(start_color='38bdf8', end_color='38bdf8', fill_type='solid')
    relleno_alterno = PatternFill(start_color='f1f5f9', end_color='f1f5f9', fill_type='solid')
    fuente_titulo = Font(name='Calibri', bold=True, color='38bdf8', size=16)
    fuente_sub = Font(name='Calibri', color='64748b', size=9, italic=True)
    fuente_kpi_et = Font(name='Calibri', color='94a3b8', size=9)
    fuente_kpi_val = Font(name='Calibri', bold=True, color='38bdf8', size=14)
    fuente_encabezado = Font(name='Calibri', bold=True, color='38bdf8', size=11)
    fuente_cuerpo = Font(name='Calibri', size=10)
    borde_sutil = Border(
        bottom=Side(style='thin', color='1e293b'),
        right=Side(style='thin', color='1e293b'),
    )
    borde_tabla = Border(
        bottom=Side(style='thin', color='cbd5e1'),
        right=Side(style='thin', color='cbd5e1'),
    )

    hoja_dash.sheet_view.showGridLines = False
    hoja_datos.sheet_view.showGridLines = False

    for col in range(1, 30):
        hoja_dash.column_dimensions[get_column_letter(col)].width = 18
    for fila in range(1, 80):
        hoja_dash.row_dimensions[fila].height = 20

    hoja_dash.merge_cells('B2:K2')
    celda_titulo = hoja_dash['B2']
    celda_titulo.value = reporte.nombre
    celda_titulo.font = fuente_titulo
    celda_titulo.fill = relleno_oscuro
    celda_titulo.alignment = Alignment(horizontal='left', vertical='center')
    hoja_dash.row_dimensions[2].height = 36

    hoja_dash.merge_cells('B3:K3')
    celda_sub = hoja_dash['B3']
    celda_sub.value = f"Tipo: {reporte.get_tipo_display()}  |  Generado: {date.today().strftime('%d/%m/%Y')}  |  Registros: {resultado['cantidad']}"
    celda_sub.font = fuente_sub
    celda_sub.fill = relleno_oscuro
    celda_sub.alignment = Alignment(horizontal='left', vertical='center')
    hoja_dash.row_dimensions[3].height = 18

    for col in range(2, 14):
        for fila in range(2, 4):
            hoja_dash.cell(row=fila, column=col).fill = relleno_oscuro

    totales = resultado.get('totales', {})
    fila_kpi = 5
    col_kpi = 2
    for etiqueta, valor in totales.items():
        hoja_dash.cell(row=fila_kpi, column=col_kpi, value=str(etiqueta)).font = fuente_kpi_et
        hoja_dash.cell(row=fila_kpi, column=col_kpi).fill = relleno_medio
        hoja_dash.cell(row=fila_kpi, column=col_kpi).alignment = Alignment(horizontal='left', vertical='center')
        hoja_dash.row_dimensions[fila_kpi].height = 16

        hoja_dash.cell(row=fila_kpi + 1, column=col_kpi, value=str(valor)).font = fuente_kpi_val
        hoja_dash.cell(row=fila_kpi + 1, column=col_kpi).fill = relleno_medio
        hoja_dash.cell(row=fila_kpi + 1, column=col_kpi).alignment = Alignment(horizontal='left', vertical='center')
        hoja_dash.row_dimensions[fila_kpi + 1].height = 28

        hoja_dash.cell(row=fila_kpi + 2, column=col_kpi).fill = relleno_oscuro
        hoja_dash.row_dimensions[fila_kpi + 2].height = 8

        col_kpi += 2
        if col_kpi > 12:
            col_kpi = 2
            fila_kpi += 3

    fila_graficas = fila_kpi + 4
    col_grafica = 2
    graficas = resultado.get('graficas', {})
    for clave, datos in graficas.items():
        etqs = datos.get('etiquetas', [])
        vals = datos.get('valores', [])
        tipo = datos.get('tipo', 'bar')
        titulo_g = datos.get('titulo', clave)
        if not etqs or not vals:
            continue

        if tipo == 'line':
            buf = _grafica_linea(etqs, vals, titulo_g, ancho=7.0, alto=3.0)
        elif tipo == 'doughnut':
            buf = _grafica_dona(etqs, vals, titulo_g, ancho=4.5, alto=3.5)
        else:
            buf = _grafica_barras(etqs, vals, titulo_g, ancho=7.0, alto=3.0)

        img_xl = XLImage(buf)
        img_xl.width = 420
        img_xl.height = 240
        celda_anclaje = get_column_letter(col_grafica) + str(fila_graficas)
        hoja_dash.add_image(img_xl, celda_anclaje)

        col_grafica += 7
        if col_grafica > 16:
            col_grafica = 2
            fila_graficas += 14

    filas = resultado.get('filas', [])
    etiquetas = resultado.get('etiquetas', {})
    if filas and etiquetas:
        claves = list(etiquetas.keys())
        encabezados = list(etiquetas.values())

        hoja_datos.merge_cells('A1:E1')
        hoja_datos['A1'] = reporte.nombre
        hoja_datos['A1'].font = Font(name='Calibri', bold=True, color='38bdf8', size=14)
        hoja_datos['A1'].fill = relleno_oscuro
        hoja_datos['A1'].alignment = Alignment(horizontal='left', vertical='center')
        hoja_datos.row_dimensions[1].height = 28

        hoja_datos['A2'] = f"Generado: {date.today().strftime('%d/%m/%Y')} — {resultado['cantidad']} registros"
        hoja_datos['A2'].font = fuente_sub
        hoja_datos['A2'].fill = relleno_oscuro
        hoja_datos.row_dimensions[2].height = 16

        fila_enc = 4
        for idx_col, enc in enumerate(encabezados, start=1):
            celda = hoja_datos.cell(row=fila_enc, column=idx_col, value=enc)
            celda.fill = relleno_medio
            celda.font = fuente_encabezado
            celda.alignment = Alignment(horizontal='center', vertical='center')
            celda.border = borde_tabla
        hoja_datos.row_dimensions[fila_enc].height = 24

        for idx_fila, fila in enumerate(filas, start=fila_enc + 1):
            relleno_fila = relleno_alterno if idx_fila % 2 == 0 else None
            for idx_col, clave in enumerate(claves, start=1):
                valor = fila.get(clave, '')
                es_num = isinstance(valor, (int, float))
                celda = hoja_datos.cell(row=idx_fila, column=idx_col,
                                        value=valor if es_num else str(valor))
                celda.font = fuente_cuerpo
                celda.alignment = Alignment(vertical='center')
                celda.border = borde_tabla
                if relleno_fila:
                    celda.fill = relleno_fila
            hoja_datos.row_dimensions[idx_fila].height = 18

        for idx_col in range(1, len(encabezados) + 1):
            ancho = max(
                (len(str(hoja_datos.cell(row=r, column=idx_col).value or ''))
                 for r in range(fila_enc, fila_enc + len(filas) + 1)),
                default=10,
            )
            hoja_datos.column_dimensions[get_column_letter(idx_col)].width = min(ancho + 4, 40)

        hoja_datos.freeze_panes = hoja_datos.cell(row=fila_enc + 1, column=1)

    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    respuesta = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    nombre_archivo = reporte.nombre.replace(' ', '_').lower()
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre_archivo}_{date.today()}.xlsx"'
    return respuesta
